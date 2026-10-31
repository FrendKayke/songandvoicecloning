"""Convertir des fichiers en PDF, en local (onglet « Outils → Convertir en PDF »).

Chaque format passe par le meilleur outil présent sur le PC, sans rien envoyer sur Internet :
  - PDF : copié tel quel ;
  - images (PNG, JPG, WebP, GIF, BMP, TIFF multipage, HEIC des iPhone) : Pillow, une page par image ;
  - documents Office (Word, Excel, PowerPoint, OpenDocument, RTF) : Microsoft Office s'il est installé (exportation
    PDF de Word / Excel / PowerPoint, pilotés par PowerShell : mise en page identique), sinon LibreOffice s'il est
    installé (soffice --convert-to pdf), sinon une conversion approchée (docx : mammoth ; xlsx : openpyxl ; pptx :
    python-pptx) mise en page en HTML puis imprimée par le navigateur ;
  - texte, Markdown, CSV, JSON, code, pages web (HTML), SVG : mis en page en HTML puis imprimés en PDF par le
    navigateur en mode sans fenêtre (Edge est présent sur tout Windows 11 ; Chrome ou Chromium sinon) ;
  - autre format : LibreOffice s'il est installé (il lit beaucoup de formats), sinon refus avec la marche à suivre.
Plusieurs fichiers : un PDF par fichier et, sur demande, un PDF unique dans l'ordre (pypdf). Sortie :
data/pdf/<horodatage>_<nom>/.
"""
import base64
import csv
import html
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import gradio as gr

from . import config as cfg
from .outils import nouveau_dossier, chemin_fichier

IMAGES = {".png", ".jpg", ".jpeg", ".jfif", ".webp", ".gif", ".bmp", ".tif", ".tiff", ".heic", ".heif", ".avif",
          ".ico"}
WORD = {".doc", ".docx", ".docm", ".dot", ".dotx", ".rtf", ".odt"}
EXCEL = {".xls", ".xlsx", ".xlsm", ".xlsb", ".ods"}
POWERPOINT = {".ppt", ".pptx", ".pptm", ".pps", ".ppsx", ".odp"}
PAGES_WEB = {".html", ".htm", ".svg", ".xhtml", ".mht", ".mhtml"}
TABLEAUX_TEXTE = {".csv", ".tsv"}
MARKDOWN = {".md", ".markdown"}
TEXTE = {".txt", ".text", ".log", ".json", ".xml", ".yaml", ".yml", ".ini", ".cfg", ".conf", ".toml", ".py", ".js",
         ".ts", ".css", ".java", ".c", ".cpp", ".h", ".cs", ".php", ".rb", ".go", ".rs", ".sql", ".sh", ".bat",
         ".ps1", ".srt", ".vtt", ".tex", ".lua", ".gd", ".nfo"}
DELAI = 300  # s : une conversion qui dépasse est abandonnée (document bloqué par une boîte de dialogue…)

STYLE = """
@page { size: A4%s; margin: 15mm; }
body { font-family: "Segoe UI", Arial, "Liberation Sans", sans-serif; font-size: 11pt; line-height: 1.45;
       color: #111; }
pre { white-space: pre-wrap; overflow-wrap: anywhere; font-family: Consolas, "DejaVu Sans Mono", monospace;
      font-size: 9.5pt; margin: 0; }
table { border-collapse: collapse; margin-bottom: 1em; }
td, th { border: 1px solid #999; padding: 2px 6px; font-size: 9pt; vertical-align: top; }
th { background: #eee; }
img { max-width: 100%%; height: auto; }
h2.feuille { font-size: 13pt; margin: 0 0 6px; }
.page { page-break-after: always; }
.page:last-child { page-break-after: auto; }
"""

# Exportation PDF par Microsoft Office (constantes de l'API COM : wdExportFormatPDF = 17, xlTypePDF = 0,
# ppSaveAsPDF = 32). Écrit en UTF-8 avec BOM : Windows PowerShell 5.1 lit sinon les accents de travers.
SCRIPT_OFFICE = r"""param([string]$Type, [string]$Entree, [string]$Sortie)
$ErrorActionPreference = 'Stop'
$app = $null
try {
    switch ($Type) {
        'word' {
            $app = New-Object -ComObject Word.Application
            $app.Visible = $false
            $app.DisplayAlerts = 0
            $doc = $app.Documents.Open($Entree, $false, $true, $false)
            $doc.ExportAsFixedFormat($Sortie, 17)
            $doc.Close(0)
        }
        'excel' {
            $app = New-Object -ComObject Excel.Application
            $app.Visible = $false
            $app.DisplayAlerts = $false
            $classeur = $app.Workbooks.Open($Entree, 0, $true)
            $classeur.ExportAsFixedFormat(0, $Sortie)
            $classeur.Close($false)
        }
        'powerpoint' {
            $app = New-Object -ComObject PowerPoint.Application
            $presentation = $app.Presentations.Open($Entree, -1, 0, 0)
            $presentation.SaveAs($Sortie, 32)
            $presentation.Close()
        }
    }
} finally {
    if ($app) {
        $app.Quit()
        [void][Runtime.InteropServices.Marshal]::ReleaseComObject($app)
    }
}
"""
APPLIS_OFFICE = {"word": "Word.Application", "excel": "Excel.Application", "powerpoint": "PowerPoint.Application"}


# --- Outils présents sur le PC ---------------------------------------------------------------------------------
def navigateur():
    """Edge, Chrome ou Chromium (impression en PDF sans fenêtre). STUDIOVOIX_NAVIGATEUR pour en imposer un."""
    impose = os.environ.get("STUDIOVOIX_NAVIGATEUR")
    if impose:
        return impose if Path(impose).is_file() else None
    candidats = []
    for racine in (os.environ.get("PROGRAMFILES(X86)"), os.environ.get("PROGRAMFILES"), os.environ.get("LOCALAPPDATA")):
        if racine:
            candidats += [Path(racine) / "Microsoft" / "Edge" / "Application" / "msedge.exe",
                          Path(racine) / "Google" / "Chrome" / "Application" / "chrome.exe"]
    for c in candidats:
        if c.is_file():
            return str(c)
    for nom in ("msedge", "google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "chrome"):
        trouve = shutil.which(nom)
        if trouve:
            return trouve
    return None


def libreoffice():
    """soffice de LibreOffice s'il est installé (STUDIOVOIX_LIBREOFFICE pour en imposer un)."""
    impose = os.environ.get("STUDIOVOIX_LIBREOFFICE")
    if impose is not None:
        return impose if impose and Path(impose).is_file() else None
    for racine in (os.environ.get("PROGRAMFILES"), os.environ.get("PROGRAMFILES(X86)")):
        if racine:
            c = Path(racine) / "LibreOffice" / "program" / "soffice.exe"
            if c.is_file():
                return str(c)
    return shutil.which("soffice") or shutil.which("libreoffice")


def office(type_):
    """L'application Office (« word », « excel », « powerpoint ») est-elle installée ? (Windows seulement)"""
    if sys.platform != "win32" or os.environ.get("STUDIOVOIX_SANS_OFFICE"):
        return False
    try:
        import winreg

        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, APPLIS_OFFICE[type_] + r"\CLSID"))
        return True
    except OSError:
        return False


def outils_disponibles():
    """Ligne d'état de l'onglet : quels outils ont été trouvés."""
    def etat(ok):
        return "✅" if ok else "❌"

    offices = [n for n in ("word", "excel", "powerpoint") if office(n)]
    return (f"Outils trouvés sur ce PC : Microsoft Office {etat(offices)}"
            + (f" ({', '.join(n.capitalize() for n in offices)})" if offices else "")
            + f" · LibreOffice {etat(libreoffice())} · navigateur (Edge / Chrome) {etat(navigateur())}. "
            "Sans Office ni LibreOffice, les fichiers Word, Excel et PowerPoint récents (.docx, .xlsx, .pptx) sont "
            "convertis avec une mise en page approchée ; installe LibreOffice (gratuit) pour une conversion fidèle "
            "et pour les anciens formats (.doc, .xls, .ppt).")


# --- Conversion d'un fichier -----------------------------------------------------------------------------------
def en_pdf(source, destination):
    """Convertit un fichier en PDF. Renvoie (méthode utilisée, remarque ou "")."""
    source, destination = Path(source), Path(destination)
    ext = source.suffix.lower()
    if ext == ".pdf":
        shutil.copy(source, destination)
        return "copié", ""
    if ext in IMAGES:
        return _images(source, destination), ""
    type_office = "word" if ext in WORD else "excel" if ext in EXCEL else "powerpoint" if ext in POWERPOINT else None
    erreurs = []
    if type_office:
        if office(type_office):
            try:
                _par_office(type_office, source, destination)
                return f"Microsoft {type_office.capitalize()}", ""
            except Exception as e:  # noqa: BLE001 - LibreOffice ou la conversion approchée prennent le relais
                erreurs.append(f"Office : {e}")
        if libreoffice():
            try:
                _par_libreoffice(source, destination)
                return "LibreOffice", ""
            except Exception as e:  # noqa: BLE001
                erreurs.append(f"LibreOffice : {e}")
        approche = {".docx": _docx_html, ".docm": _docx_html, ".xlsx": _xlsx_html, ".xlsm": _xlsx_html,
                    ".pptx": _pptx_html, ".pptm": _pptx_html}.get(ext)
        if approche:
            try:
                _imprimer(approche(source), destination)
            except gr.Error as e:  # les échecs d'Office / LibreOffice expliquent souvent le reste
                raise gr.Error(f"{source.name} : {getattr(e, 'message', e)}"
                               + (f" ({' ; '.join(erreurs)})" if erreurs else "")) from e
            return ("conversion approchée", "mise en page approchée (ni Office ni LibreOffice sur ce PC)"
                    + (f" ; {' ; '.join(erreurs)}" if erreurs else ""))
        raise gr.Error(f"{source.name} : il faut Microsoft Office ou LibreOffice (gratuit) pour ce format."
                       + (f" ({' ; '.join(erreurs)})" if erreurs else ""))
    if ext in PAGES_WEB:
        _imprimer_url(source.resolve().as_uri(), destination)
        return "navigateur", ""
    if ext in TABLEAUX_TEXTE or ext in MARKDOWN or ext in TEXTE:
        texte = _lire_texte(source)
        if ext in MARKDOWN:
            from markdown_it import MarkdownIt

            corps = MarkdownIt("commonmark").enable("table").render(texte)
        elif ext in TABLEAUX_TEXTE:
            corps = _table_html(_lignes_csv(texte, "\t" if ext == ".tsv" else None))
        else:
            corps = f"<pre>{html.escape(texte)}</pre>"
        _imprimer(_page(source.name, corps, paysage=ext in TABLEAUX_TEXTE), destination)
        return "navigateur", ""
    if libreoffice():  # format inconnu : LibreOffice en lit beaucoup (dessins, anciens traitements de texte…)
        _par_libreoffice(source, destination)
        return "LibreOffice", ""
    if _ressemble_a_du_texte(source):
        _imprimer(_page(source.name, f"<pre>{html.escape(_lire_texte(source))}</pre>"), destination)
        return "navigateur", "lu comme du texte"
    raise gr.Error(f"{source.name} : format {ext or 'sans extension'} non pris en charge. Installe LibreOffice "
                   "(gratuit) pour convertir davantage de formats.")


def _images(source, destination):
    from PIL import Image, ImageOps, ImageSequence

    try:
        from pillow_heif import register_heif_opener  # photos HEIC des iPhone

        register_heif_opener()
    except ImportError:
        pass
    pages = []
    with Image.open(source) as im:
        trames = list(ImageSequence.Iterator(im)) if source.suffix.lower() in (".tif", ".tiff") else [im]
        for trame in trames:
            t = ImageOps.exif_transpose(trame)
            if t.mode in ("RGBA", "LA", "P"):  # transparence posée sur du blanc (le PDF image n'a pas d'alpha)
                t = t.convert("RGBA")
                fond = Image.new("RGB", t.size, "white")
                fond.paste(t, mask=t.getchannel("A"))
                t = fond
            pages.append(t.convert("RGB"))
    resolution = 150.0
    pages[0].save(destination, "PDF", save_all=True, append_images=pages[1:], resolution=resolution)
    return "image"


def _par_office(type_, source, destination):
    with tempfile.TemporaryDirectory() as tmp:
        script = Path(tmp) / "office_pdf.ps1"
        script.write_text(SCRIPT_OFFICE, encoding="utf-8-sig")
        r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File",
                            str(script), "-Type", type_, "-Entree", str(source.resolve()),
                            "-Sortie", str(destination.resolve())],
                           capture_output=True, text=True, timeout=DELAI, creationflags=_sans_fenetre())
    if r.returncode != 0 or not destination.exists():
        lignes = (r.stderr or r.stdout or "").strip().splitlines()
        raise RuntimeError(lignes[-1][:300] if lignes else "aucun PDF produit")


def _par_libreoffice(source, destination):
    with tempfile.TemporaryDirectory() as tmp:
        profil = Path(tmp) / "profil"  # profil à part : marche même si LibreOffice est déjà ouvert
        r = subprocess.run([libreoffice(), f"-env:UserInstallation={profil.as_uri()}", "--headless", "--norestore",
                            "--convert-to", "pdf", "--outdir", tmp, str(source.resolve())],
                           capture_output=True, text=True, timeout=DELAI, creationflags=_sans_fenetre())
        produit = Path(tmp) / (source.stem + ".pdf")
        if not produit.exists():
            raise RuntimeError((r.stderr or r.stdout or "aucun PDF produit").strip()[-300:])
        shutil.move(str(produit), destination)


def _imprimer(page_html, destination):
    """Page HTML → PDF par le navigateur sans fenêtre."""
    with tempfile.TemporaryDirectory() as tmp:
        fichier = Path(tmp) / "page.html"
        fichier.write_text(page_html, encoding="utf-8")
        _imprimer_url(fichier.as_uri(), destination)


def _imprimer_url(url, destination):
    nav = navigateur()
    if not nav:
        raise gr.Error("Aucun navigateur trouvé pour imprimer en PDF (Edge, présent sur Windows 11, ou Chrome).")
    destination = Path(destination).resolve()
    with tempfile.TemporaryDirectory() as profil:  # profil à part : marche même si Edge est déjà ouvert
        commande = [nav, "--headless", "--disable-gpu", "--no-first-run", "--no-default-browser-check",
                    "--disable-extensions", "--no-pdf-header-footer", f"--user-data-dir={profil}",
                    f"--print-to-pdf={destination}", url]
        if os.environ.get("STUDIOVOIX_SANS_BAC_A_SABLE") or (hasattr(os, "geteuid") and os.geteuid() == 0):
            # Linux (tests, vérification automatique) : Chromium refuse son bac à sable en administrateur ou quand
            # AppArmor interdit les espaces de noms ; jamais nécessaire sous Windows
            commande.insert(1, "--no-sandbox")
        try:
            r = subprocess.run(commande, capture_output=True, text=True, timeout=DELAI, creationflags=_sans_fenetre())
        except subprocess.TimeoutExpired as e:
            raise gr.Error("Le navigateur n'a pas fini l'impression en PDF à temps.") from e
    for _ in range(20):  # le fichier peut apparaître un instant après la fin du processus
        if destination.exists() and destination.stat().st_size > 0:
            return
        time.sleep(0.1)
    raise gr.Error(f"Le navigateur n'a pas produit de PDF : {(r.stderr or r.stdout or '').strip()[-300:]}")


def _sans_fenetre():
    return getattr(subprocess, "CREATE_NO_WINDOW", 0)


# --- Mise en page HTML ------------------------------------------------------------------------------------------
def _page(titre, corps, paysage=False):
    style = STYLE % (" landscape" if paysage else "",)
    return (f'<!doctype html><html lang="fr"><head><meta charset="utf-8"><title>{html.escape(titre)}</title>'
            f"<style>{style}</style></head><body>{corps}</body></html>")


def _lire_texte(source):
    brut = Path(source).read_bytes()
    for codage in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return brut.decode(codage)
        except UnicodeDecodeError:
            continue
    return brut.decode("utf-8", errors="replace")


def _ressemble_a_du_texte(source):
    debut = Path(source).read_bytes()[:4096]
    return bool(debut) and b"\x00" not in debut


def _lignes_csv(texte, separateur=None):
    if separateur is None:
        try:
            separateur = csv.Sniffer().sniff(texte[:20000], delimiters=";,\t|").delimiter
        except csv.Error:
            separateur = ";" if texte.count(";") > texte.count(",") else ","
    return list(csv.reader(io.StringIO(texte), delimiter=separateur))


def _table_html(lignes, entete=True):
    if not lignes:
        return "<p>(vide)</p>"
    largeur = max(len(l_) for l_ in lignes)
    rangs = []
    for i, l_ in enumerate(lignes):
        balise = "th" if entete and i == 0 else "td"
        cellules = "".join(f"<{balise}>{html.escape(str(c))}</{balise}>" for c in list(l_) + [""] * (largeur - len(l_)))
        rangs.append(f"<tr>{cellules}</tr>")
    return f"<table>{''.join(rangs)}</table>"


def _docx_html(source):
    import mammoth

    with open(source, "rb") as f:
        resultat = mammoth.convert_to_html(f)  # images incluses dans la page (data:)
    return _page(source.name, resultat.value)


def _xlsx_html(source):
    import openpyxl

    classeur = openpyxl.load_workbook(source, read_only=True, data_only=True)
    parties = []
    for feuille in classeur.worksheets:
        lignes = [["" if v is None else v for v in ligne] for ligne in feuille.iter_rows(values_only=True)]
        while lignes and not any(str(c).strip() for c in lignes[-1]):
            lignes.pop()
        if lignes:
            parties.append(f'<div class="page"><h2 class="feuille">{html.escape(feuille.title)}</h2>'
                           f"{_table_html(lignes, entete=False)}</div>")
    classeur.close()
    return _page(source.name, "".join(parties) or "<p>(classeur vide)</p>", paysage=True)


def _pptx_html(source):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    pres = Presentation(source)
    diapos = []
    for n, diapo in enumerate(pres.slides, 1):
        morceaux = []
        for forme in diapo.shapes:
            if forme.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = forme.image
                donnees = base64.b64encode(image.blob).decode()
                morceaux.append(f'<img src="data:{image.content_type};base64,{donnees}">')
            elif forme.has_text_frame and forme.text_frame.text.strip():
                titre = diapo.shapes.title is not None and forme.shape_id == diapo.shapes.title.shape_id
                texte = html.escape(forme.text_frame.text).replace("\n", "<br>")
                morceaux.append(f"<h2>{texte}</h2>" if titre else f"<p>{texte}</p>")
        diapos.append(f'<div class="page"><p style="color:#888;font-size:9pt">Diapositive {n}</p>'
                      f"{''.join(morceaux)}</div>")
    return _page(source.name, "".join(diapos) or "<p>(présentation vide)</p>", paysage=True)


# --- Onglet -----------------------------------------------------------------------------------------------------
def _nom_libre(dossier, nom):
    """nom.pdf, ou nom_2.pdf… s'il est déjà pris (deux fichiers « rapport.docx » et « rapport.xlsx »)."""
    candidat, k = dossier / f"{nom}.pdf", 1
    while candidat.exists():
        k += 1
        candidat = dossier / f"{nom}_{k}.pdf"
    return candidat


def convertir(fichiers, fusionner=False, nom="", progress=gr.Progress()):
    """Convertit chaque fichier en PDF (et, si demandé, les réunit dans un seul PDF, dans l'ordre).
    Renvoie (message, fichiers PDF, dossier)."""
    chemins = [Path(str(chemin_fichier(f))) for f in (fichiers or []) if f]
    if not chemins:
        raise gr.Error("Ajoute un ou plusieurs fichiers à convertir.")
    nom = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", (nom or "").strip()).strip(". ") or "documents"
    dossier = nouveau_dossier(cfg.PDF_DIR, nom)
    debut, faits, lignes, erreurs = time.monotonic(), [], [], []
    for i, source in enumerate(chemins):
        progress(i / len(chemins), desc=f"Conversion {i + 1}/{len(chemins)} : {source.name}")
        destination = _nom_libre(dossier, source.stem)
        try:
            methode, remarque = en_pdf(source, destination)
        except subprocess.TimeoutExpired:
            erreurs.append(f"{source.name} : la conversion a dépassé {DELAI // 60} min (document bloqué ?).")
            continue
        except gr.Error as e:
            erreurs.append(getattr(e, "message", str(e)))
            continue
        except Exception as e:  # noqa: BLE001 - fichier abîmé, protégé par mot de passe…
            erreurs.append(f"{source.name} : {e}")
            continue
        faits.append(destination)
        lignes.append(f"- {source.name} → {destination.name} ({methode}{' ; ' + remarque if remarque else ''})")
    if not faits:
        raise gr.Error("Aucun fichier n'a pu être converti :\n" + "\n".join(erreurs))
    sorties = list(faits)
    if fusionner and len(faits) > 1:
        from pypdf import PdfWriter

        progress(0.95, desc="Assemblage en un seul PDF…")
        ensemble = _nom_libre(dossier, f"{nom}_complet")
        redacteur = PdfWriter()
        for f in faits:
            redacteur.append(str(f))
        with open(ensemble, "wb") as sortie:
            redacteur.write(sortie)
        sorties.insert(0, ensemble)
        lignes.insert(0, f"- **{ensemble.name}** : les {len(faits)} PDF réunis dans l'ordre")
    msg = (f"✅ {len(faits)}/{len(chemins)} fichier(s) converti(s) en {time.monotonic() - debut:.0f} s, dans "
           f"{dossier}.\n\n" + "\n".join(lignes))
    if erreurs:
        msg += "\n\n⚠️ " + "\n\n⚠️ ".join(erreurs)
    return msg, [str(s) for s in sorties], str(dossier)

