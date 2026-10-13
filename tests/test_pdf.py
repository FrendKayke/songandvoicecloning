"""Convertir en PDF : images (Pillow), texte / Markdown / CSV / HTML (navigateur sans fenêtre), Word / Excel /
PowerPoint (LibreOffice, ou conversion approchée), assemblage en un seul PDF."""
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import gradio as gr
import pytest

from conftest import no_progress
from studiovoix import config as cfg
from studiovoix import pdf

CHROMIUM = Path("/opt/pw-browsers/chromium-1194/chrome-linux/chrome")  # navigateur de l'environnement de test


@pytest.fixture
def nav(monkeypatch):
    if CHROMIUM.is_file():
        monkeypatch.setenv("STUDIOVOIX_NAVIGATEUR", str(CHROMIUM))
    if sys.platform.startswith("linux"):
        monkeypatch.setenv("STUDIOVOIX_SANS_BAC_A_SABLE", "1")
    if not pdf.navigateur():
        pytest.skip("aucun navigateur (Edge, Chrome, Chromium) pour imprimer en PDF")


def _texte(chemin):
    from pypdf import PdfReader

    return " ".join(p.extract_text() or "" for p in PdfReader(str(chemin)).pages)


def _pages(chemin):
    from pypdf import PdfReader

    return len(PdfReader(str(chemin)).pages)


def _docx(chemin, texte):
    """Document Word minimal (les trois parties obligatoires)."""
    with zipfile.ZipFile(chemin, "w") as z:
        z.writestr("[Content_Types].xml",
                   '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/'
                   '2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-'
                   'package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override '
                   'PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.'
                   'wordprocessingml.document.main+xml"/></Types>')
        z.writestr("_rels/.rels",
                   '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/'
                   'package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
                   'officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr("word/document.xml",
                   '<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/'
                   'wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>' + texte + '</w:t></w:r></w:p></w:body>'
                   '</w:document>')
    return chemin


def _office(dossier):
    import openpyxl
    from pptx import Presentation
    from pptx.util import Inches

    docx = _docx(dossier / "lettre.docx", "Bonjour Amélie, voici le devis")
    classeur = openpyxl.Workbook()
    classeur.active.title = "Budget"
    classeur.active.append(["Poste", "Montant"])
    classeur.active.append(["Décor", 1250])
    xlsx = dossier / "budget.xlsx"
    classeur.save(xlsx)
    pres = Presentation()
    diapo = pres.slides.add_slide(pres.slide_layouts[1])
    diapo.shapes.title.text = "Présentation du jeu"
    diapo.placeholders[1].text = "Trois niveaux, un dragon"
    pptx = dossier / "jeu.pptx"
    pres.save(pptx)
    return docx, xlsx, pptx


def test_images_en_pdf(env):
    from PIL import Image

    png = env / "logo.png"
    Image.new("RGBA", (200, 100), (255, 0, 0, 0)).save(png)
    jpg = env / "photo.jpg"
    Image.new("RGB", (300, 200), "blue").save(jpg)
    tif = env / "scan.tif"
    Image.new("RGB", (100, 100), "white").save(tif, save_all=True, append_images=[Image.new("RGB", (100, 100))] * 2)
    msg, sorties, dossier = pdf.convertir([str(png), str(jpg), str(tif)], True, "photos", progress=no_progress)
    noms = [Path(s).name for s in sorties]
    assert noms == ["photos_complet.pdf", "logo.pdf", "photo.pdf", "scan.pdf"] and "3/3" in msg
    assert _pages(sorties[3]) == 3 and _pages(sorties[0]) == 5  # TIFF de 3 pages ; tout réuni dans l'ordre
    assert Path(dossier).parent == cfg.PDF_DIR and Path(dossier).name.endswith("_photos")
    try:  # photo HEIC d'iPhone
        from pillow_heif import register_heif_opener

        register_heif_opener()
        heic = env / "iphone.heic"
        Image.new("RGB", (64, 48), "green").save(heic)
    except Exception:  # noqa: BLE001 - encodeur HEIF absent : seule la lecture compte chez l'utilisateur
        return
    assert pdf.en_pdf(heic, env / "iphone.pdf")[0] == "image" and _pages(env / "iphone.pdf") == 1


def test_texte_markdown_csv_html(env, nav):
    txt = env / "notes.txt"
    txt.write_bytes("Liste des courses : pâtes, crème\n".encode("cp1252"))  # ancien fichier Windows
    md = env / "lisez-moi.md"
    md.write_text("# Titre du projet\n\n| Nom | Rôle |\n|---|---|\n| Léa | héroïne |\n", encoding="utf-8")
    tableau = env / "scores.csv"
    tableau.write_text("joueur;score\nZoé;42\n", encoding="utf-8")
    page = env / "page.html"
    page.write_text("<html><body><h1>Bienvenue à bord</h1></body></html>", encoding="utf-8")
    msg, sorties, _ = pdf.convertir([str(txt), str(md), str(tableau), str(page)], False, "", progress=no_progress)
    assert "4/4" in msg and len(sorties) == 4
    textes = [_texte(s) for s in sorties]
    assert "crème" in textes[0] and "Titre du projet" in textes[1] and "héroïne" in textes[1]
    assert "Zoé" in textes[2] and "42" in textes[2] and "Bienvenue à bord" in textes[3]


def test_office_sans_office_ni_libreoffice(env, nav, monkeypatch):
    """Conversion approchée : docx (mammoth), xlsx (openpyxl), pptx (python-pptx), mis en page par le navigateur."""
    monkeypatch.setenv("STUDIOVOIX_LIBREOFFICE", "")
    monkeypatch.setenv("STUDIOVOIX_SANS_OFFICE", "1")
    assert pdf.libreoffice() is None and "LibreOffice ❌" in pdf.outils_disponibles()
    docx, xlsx, pptx = _office(env)
    msg, sorties, _ = pdf.convertir([str(docx), str(xlsx), str(pptx)], False, "", progress=no_progress)
    assert "3/3" in msg and msg.count("conversion approchée") == 3
    textes = [_texte(s) for s in sorties]
    assert "Bonjour Amélie" in textes[0] and "Décor" in textes[1] and "1250" in textes[1]
    assert "Présentation du jeu" in textes[2] and "un dragon" in textes[2]
    ancien = env / "vieux.doc"
    ancien.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 100)
    with pytest.raises(gr.Error, match="LibreOffice"):
        pdf.convertir([str(ancien)], False, "", progress=no_progress)
    inconnu = env / "modele.blend"
    inconnu.write_bytes(b"BLENDER\x00\x01")
    msg, sorties, _ = pdf.convertir([str(inconnu), str(docx)], False, "", progress=no_progress)
    assert "1/2" in msg and "non pris en charge" in msg


@pytest.mark.skipif(not pdf.libreoffice(), reason="LibreOffice absent")
def test_office_par_libreoffice(env):
    docx, xlsx, pptx = _office(env)
    msg, sorties, _ = pdf.convertir([str(docx), str(xlsx), str(pptx)], True, "dossier", progress=no_progress)
    assert "3/3" in msg and msg.count("LibreOffice") == 3
    assert "Bonjour Amélie" in _texte(sorties[1]) and "Présentation du jeu" in _texte(sorties[3])
    assert _pages(sorties[0]) == sum(_pages(s) for s in sorties[1:])


def test_noms_et_erreurs(env, monkeypatch):
    from PIL import Image

    a = env / "a"
    b = env / "b"
    a.mkdir()
    b.mkdir()
    Image.new("RGB", (10, 10)).save(a / "rapport.png")
    shutil.copy(a / "rapport.png", b / "rapport.jpg")
    msg, sorties, _ = pdf.convertir([str(a / "rapport.png"), str(b / "rapport.jpg")], False, "x/y:z",
                                    progress=no_progress)
    assert [Path(s).name for s in sorties] == ["rapport.pdf", "rapport_2.pdf"] and Path(sorties[0]).parent.name.endswith("x_y_z")
    copie = pdf.en_pdf(sorties[0], env / "copie.pdf")
    assert copie[0] == "copié" and _pages(env / "copie.pdf") == 1
    with pytest.raises(gr.Error, match="Ajoute"):
        pdf.convertir(None, False, "", progress=no_progress)
    monkeypatch.setenv("STUDIOVOIX_NAVIGATEUR", str(env / "absent.exe"))
    (env / "n.txt").write_text("x", encoding="utf-8")
    with pytest.raises(gr.Error, match="Aucun fichier.*navigateur"):
        pdf.convertir([str(env / "n.txt")], False, "", progress=no_progress)


@pytest.mark.skipif(sys.platform != "win32", reason="Windows PowerShell seulement")
def test_script_office_valide_pour_powershell(tmp_path):
    """Le script d'exportation par Office (non exécutable sans Office) est au moins valide pour PowerShell 5.1."""
    script = tmp_path / "office_pdf.ps1"
    script.write_text(pdf.SCRIPT_OFFICE, encoding="utf-8-sig")
    commande = ("$e = $null; [void][System.Management.Automation.Language.Parser]::ParseFile("
                f"'{script}', [ref]$null, [ref]$e); if ($e.Count) {{ $e | ForEach-Object {{ $_.Message }}; exit 1 }}")
    r = subprocess.run(["powershell", "-NoProfile", "-Command", commande], capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr
    assert os.environ.get("STUDIOVOIX_SANS_OFFICE") is None or not pdf.office("word")
